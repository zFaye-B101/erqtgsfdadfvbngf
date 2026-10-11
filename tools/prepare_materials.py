"""Copy assignment materials and create searchable companions, without dataset files."""
import argparse
import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader


def link(path):
    return quote(str(path).replace('\\', '/'), safe='/')


def table_md(rows):
    clean = [[str(v).replace('|', '\\|').replace('\n', '<br>') for v in row] for row in rows]
    if not clean:
        return ''
    width = max(map(len, clean))
    clean = [row + [''] * (width - len(row)) for row in clean]
    lines = ['| ' + ' | '.join(row) + ' |' for row in clean]
    lines.insert(1, '| ' + ' | '.join(['---'] * width) + ' |')
    return '\n'.join(lines)


def convert(path):
    header = f'# {path.stem}\n\nSource: [{path.name}]({link(path.name)})\n\n'
    notes = 'Machine-extracted reading copy, not a replacement for the original. Verify equations, figures and layout against the linked original.\n\n'
    body = []
    if path.suffix.lower() == '.pdf':
        reader = PdfReader(path)
        for number in range(1, len(reader.pages) + 1):
            result = subprocess.run(['pdftotext', '-f', str(number), '-l', str(number), '-layout', str(path), '-'], check=True, capture_output=True, text=True)
            page_text = result.stdout.replace('\f', '').strip()
            if len(page_text) < 40:
                raise ValueError(f'Page {number} needs OCR: {path.name}')
            body.append(f'## Page {number}\n\n```text\n{page_text}\n```')
    elif path.suffix.lower() == '.docx':
        doc = Document(path)
        for child in doc.element.body.iterchildren():
            if child.tag.endswith('}p'):
                paragraph = Paragraph(child, doc)
                if paragraph.text.strip():
                    body.append(paragraph.text)
            elif child.tag.endswith('}tbl'):
                table = Table(child, doc)
                body.append(table_md([[cell.text for cell in row.cells] for row in table.rows]))
        for section in doc.sections:
            for part in (section.header, section.footer):
                text = '\n'.join(p.text for p in part.paragraphs if p.text.strip())
                if text and text not in body:
                    body.append('### Header / footer\n\n' + text)
    elif path.suffix.lower() == '.pptx':
        prs = Presentation(path)
        assets = path.parent / (path.stem + '-assets')
        def shapes_text(shapes, slide_no):
            result = []
            for shape in shapes:
                if shape.has_text_frame and shape.text.strip():
                    result.append(shape.text)
                if shape.has_table:
                    result.append(table_md([[c.text for c in row.cells] for row in shape.table.rows]))
                if hasattr(shape, 'image'):
                    assets.mkdir(exist_ok=True)
                    name = f'slide-{slide_no:02d}-shape-{shape.shape_id}.{shape.image.ext}'
                    (assets / name).write_bytes(shape.image.blob)
                    result.append(f'![Slide {slide_no} embedded image; inspect original for layout]({link(assets.name + "/" + name)})')
                if hasattr(shape, 'shapes'):
                    result.extend(shapes_text(shape.shapes, slide_no))
            return result
        for number, slide in enumerate(prs.slides, 1):
            pieces = shapes_text(slide.shapes, number)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame and frame.text.strip():
                    pieces.append('### Speaker notes\n\n' + frame.text)
            body.append(f'## Slide {number}\n\n' + '\n\n'.join(pieces))
    else:
        return None
    output = path.with_suffix('.md')
    output.write_text(header + notes + '\n\n'.join(body) + '\n', encoding='utf-8')
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--paper', action='append', default=[], help='Repository PDF name=local source path')
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    dest = repo / 'materials'
    records = []
    for src in sorted(args.source.rglob('*')):
        if not src.is_file():
            continue
        rel = src.relative_to(args.source)
        reason = None
        if 'data' in rel.parts:
            reason = 'Dataset excluded'
        elif src.name == '.DS_Store':
            reason = 'Operating-system metadata excluded'
        elif src.stat().st_size == 0:
            reason = 'Empty placeholder excluded'
        if reason:
            records.append({'source': str(rel), 'status': reason})
            continue
        target = dest / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, target)
        status = 'Original copied'
        if src.suffix == '.ipynb':
            nb = json.loads(src.read_text(encoding='utf-8'))
            for cell in nb['cells']:
                cell.pop('attachments', None)
                if cell['cell_type'] == 'code':
                    cell['outputs'] = []
                    cell['execution_count'] = None
                cell['metadata'] = {}
            nb['metadata'].pop('widgets', None)
            target.write_text(json.dumps(nb, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
            md = ['# Data audit notebook', 'Code-only export. Stored outputs, including dataset pictures, are intentionally excluded. This is not the final model notebook.']
            for i, cell in enumerate(nb['cells'], 1):
                source = ''.join(cell.get('source', []))
                md.append(f'## Cell {i}\n\n' + (f'```python\n{source}\n```' if cell['cell_type'] == 'code' else source))
            target.with_suffix('.md').write_text('\n\n'.join(md) + '\n', encoding='utf-8')
            status = 'Code unchanged; outputs, attachments and execution metadata removed'
        elif src.suffix == '.md':
            text = re.sub(r'/Users/[^/]+/Downloads/28749213\.html', '28749213.html', target.read_text(encoding='utf-8'))
            if rel == Path('README.md') or rel.parts[0] == 'planning' or src.name == 'uci-har-scarf.md':
                text = '> Historical document. Model and submission-status statements may be outdated. See the repository README and docs/current-plan.md for the current status.\n\n' + text
                status = 'Historical banner added; local absolute path sanitized where present'
            target.write_text(text, encoding='utf-8')
        companion = convert(target)
        records.append({'source': str(rel), 'repository_path': str(target.relative_to(repo)), 'status': status,
                        'original_sha256': hashlib.sha256(src.read_bytes()).hexdigest(),
                        'uploaded_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
                        'markdown': str(companion.relative_to(repo)) if companion else None})
    papers = repo / 'papers'
    papers.mkdir(exist_ok=True)
    for item in args.paper:
        name, local_path = item.split('=', 1)
        if Path(name).name != name or not name.endswith('.pdf'):
            raise ValueError('Paper name must be a simple PDF filename')
        target = papers / name
        shutil.copy2(local_path, target)
        convert(target)
        records.append({'source': name, 'repository_path': str(target.relative_to(repo)), 'status': 'Paper original and reading copy', 'uploaded_sha256': hashlib.sha256(target.read_bytes()).hexdigest(), 'markdown': str(target.with_suffix('.md').relative_to(repo))})
    (repo / 'docs').mkdir(exist_ok=True)
    (repo / 'docs/materials-manifest.json').write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'included_sources': sum('repository_path' in r for r in records), 'excluded': [r for r in records if 'repository_path' not in r]}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
